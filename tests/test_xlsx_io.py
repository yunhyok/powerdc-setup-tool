"""`core/xlsx_io.py` -- the v0.1.4 Excel round trip of the VRM/Sink tables.

The contract under test, in the user's own words: *export both tables, edit the
values in Excel, import them back; the row order may change and must not
matter, but the net (channel) names and the column structure must not change.*

So the shape of this module is: one round trip that must be a no-op, a handful
of edits that must land (with their override flags), a row-reordered workbook
that must still match, and a tamper matrix in which every single entry has to
be a blocking error with **nothing** applied.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from fixtures import GROUND_NET, POWER_NET_A, POWER_NET_B, build_mini_spd
from powerdc_setup_tool.core.columns import SINK_COLUMNS, VRM_COLUMNS
from powerdc_setup_tool.core.session import Session, sink_key, vrm_key
from powerdc_setup_tool.core.spd_scan import scan_spd
from powerdc_setup_tool.core.xlsx_io import (
    KEY_HEADER,
    SCHEMA_VERSION,
    SHEET_META,
    SHEET_SINK,
    SHEET_VRM,
    ImportReport,
    apply_report,
    export_tables,
    import_tables,
)

VRM_A = vrm_key(POWER_NET_A)
VRM_B = vrm_key(POWER_NET_B)
SINK_A = sink_key(POWER_NET_A)
SINK_B = sink_key(POWER_NET_B)


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #


@pytest.fixture
def session(tmp_path: Path) -> Session:
    source = tmp_path / "mini.spd"
    build_mini_spd(source, style="si")
    loaded = Session()
    loaded.load(scan_spd(source))
    return loaded


@pytest.fixture
def book(session: Session, tmp_path: Path) -> Path:
    """An exported workbook, ready to be edited and imported back."""
    return export_tables(session, tmp_path / "tables.xlsx", exported_at="2026-08-19T09:00:00")


def _rows(path: Path, sheet: str) -> list[dict[str, Any]]:
    """The sheet as ``[{header: value}]`` -- what the tests edit."""
    workbook = load_workbook(path)
    worksheet = workbook[sheet]
    headers = [cell.value for cell in worksheet[1]]
    records = [
        dict(zip(headers, row))
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]
    workbook.close()
    return records


def _rewrite(
    path: Path,
    sheet: str,
    records: list[dict[str, Any]],
    *,
    headers: list[str] | None = None,
) -> None:
    """Put *records* back on *sheet*, optionally under a different column set."""
    workbook = load_workbook(path)
    position = workbook.sheetnames.index(sheet)
    if headers is None:
        headers = [cell.value for cell in workbook[sheet][1]]
    del workbook[sheet]
    worksheet = workbook.create_sheet(sheet, position)
    worksheet.append(headers)
    for record in records:
        worksheet.append([record.get(header) for header in headers])
    workbook.save(path)
    workbook.close()


def _edit(path: Path, sheet: str, row_key: str, header: str, value: Any) -> None:
    """Type *value* into one cell, addressed the way a user would see it."""
    records = _rows(path, sheet)
    matches = [record for record in records if record[KEY_HEADER] == row_key]
    assert matches, f"{row_key} not on sheet {sheet}"
    for record in matches:
        record[header] = value
    _rewrite(path, sheet, records)


def _headers(path: Path, sheet: str) -> list[str]:
    workbook = load_workbook(path)
    values = [cell.value for cell in workbook[sheet][1]]
    workbook.close()
    return values


def _meta(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path)
    values = {
        str(row[0]): row[1]
        for row in workbook[SHEET_META].iter_rows(values_only=True)
        if row and row[0] is not None
    }
    workbook.close()
    return values


def _snapshot(session: Session) -> list[tuple[Any, ...]]:
    """Every value the import could possibly touch, for "nothing was applied"."""
    return [
        (
            row.net,
            row.row_id,
            row.enabled,
            row.comp,
            row.gnet,
            row.nominal_voltage,
            getattr(row, "sense_voltage", None),
            getattr(row, "output_current", getattr(row, "current", None)),
            getattr(row, "model", None),
            getattr(row, "pf_mode", None),
            getattr(row, "pin_equal_current", None),
        )
        for row in (*session.vrm_rows, *session.sink_rows)
    ]


def _assert_blocked(session: Session, report: ImportReport, *needles: str) -> None:
    """A tamper-matrix entry: blocking error, named cell, nothing applied."""
    assert report.errors, "expected a blocking error"
    assert not report.ok
    joined = " | ".join(report.errors)
    for needle in needles:
        assert needle in joined, f"{needle!r} not in {joined!r}"
    assert report.applied_changes == []
    with pytest.raises(ValueError):
        apply_report(session, report)


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #


def test_export_writes_both_tables_and_a_hidden_meta_sheet(book: Path) -> None:
    workbook = load_workbook(book)
    assert workbook.sheetnames == [SHEET_VRM, SHEET_SINK, SHEET_META]
    assert workbook[SHEET_META].sheet_state == "hidden"
    assert workbook[SHEET_VRM].sheet_state == "visible"
    workbook.close()

    meta = _meta(book)
    assert meta["schema"] == SCHEMA_VERSION
    assert meta["source"] == "mini.spd"
    assert meta["exported_at"] == "2026-08-19T09:00:00"
    assert meta["vrm_rows"] == 2 and meta["sink_rows"] == 2


def test_export_headers_are_key_plus_the_gui_columns(book: Path) -> None:
    assert _headers(book, SHEET_VRM) == [KEY_HEADER, *(c.title for c in VRM_COLUMNS)]
    assert _headers(book, SHEET_SINK) == [KEY_HEADER, *(c.title for c in SINK_COLUMNS)]


def test_export_carries_values_reference_columns_and_formatting(
    session: Session, book: Path
) -> None:
    vrms = {record[KEY_HEADER]: record for record in _rows(book, SHEET_VRM)}
    assert set(vrms) == {VRM_A, VRM_B}
    row = vrms[VRM_A]
    assert row["Net"] == POWER_NET_A  # read-only reference column, for context
    assert row["Use"] is True
    assert row["Component"] == "LGA"
    assert row["Ground"] == GROUND_NET
    assert row["Nominal V"] == pytest.approx(0.7)
    assert row["Sense V"] == pytest.approx(0.7)
    assert row["Pos pins"] == len(session.positive_pins(session.vrm(POWER_NET_A)))
    assert row["Name"].startswith("VRM_LGA_")

    sinks = {record[KEY_HEADER]: record for record in _rows(book, SHEET_SINK)}
    assert sinks[SINK_A]["Component"] == "SITE0"
    assert (sinks[SINK_A]["Model"], sinks[SINK_A]["PFMode"]) == (2, 2)
    assert sinks[SINK_A]["PinEqualCurrent"] == 1

    workbook = load_workbook(book)
    worksheet = workbook[SHEET_VRM]
    assert worksheet.freeze_panes == "A2"  # header stays put while scrolling
    assert worksheet.cell(row=1, column=1).font.bold
    # Sheet protection is deliberately OFF: Excel forbids moving rows on a
    # protected sheet, and reordering rows is the point (see the module docstring).
    assert not worksheet.protection.sheet
    assert worksheet.cell(row=2, column=6).number_format == "0.####"
    workbook.close()


# --------------------------------------------------------------------------- #
# Round trip / edits
# --------------------------------------------------------------------------- #


def test_untouched_round_trip_changes_nothing(session: Session, book: Path) -> None:
    report = import_tables(session, book)
    assert report.errors == []
    assert report.warnings == []
    assert report.applied_changes == []
    assert report.cell_count == 0 and report.row_count == 0
    assert apply_report(session, report) == []


def test_value_edits_apply_and_raise_the_override_flags(session: Session, book: Path) -> None:
    _edit(book, SHEET_VRM, VRM_A, "Nominal V", 0.85)
    _edit(book, SHEET_VRM, VRM_A, "Output Current (A)", 12.5)
    _edit(book, SHEET_SINK, SINK_B, "Current (A)", 7)
    _edit(book, SHEET_SINK, SINK_B, "PinEqualCurrent", 0)

    report = import_tables(session, book)
    assert report.errors == []
    assert report.cell_count == 4
    assert report.row_count == 2
    assert (report.vrm_row_count, report.sink_row_count) == (1, 1)
    assert (VRM_A, "nominal_voltage", 0.7, 0.85) in report.applied_changes

    apply_report(session, report)
    vrm = session.vrm(POWER_NET_A)
    assert (vrm.nominal_voltage, vrm.output_current) == (0.85, 12.5)
    assert vrm.nominal_override and vrm.current_override
    sink = session.sink(POWER_NET_B)
    assert (sink.current, sink.pin_equal_current) == (7.0, 0)
    assert sink.current_override
    # An edited nominal must not drag the net voltage with it (design §B).
    assert session.nets[POWER_NET_A].voltage == pytest.approx(0.7)


def test_an_imported_nominal_drags_its_auto_sense_along(session: Session, book: Path) -> None:
    """Exactly what typing the same value into the GUI cell does (design §B).

    `sense_voltage` follows its row's nominal while it is still auto, so a
    workbook that moves *only* the nominal moves the sense with it -- the sense
    cell of that workbook is stale the moment the user edits the nominal, and
    honouring it instead would quietly pin a cell the user never touched.
    """
    _edit(book, SHEET_VRM, VRM_A, "Nominal V", 0.85)

    report = import_tables(session, book)
    assert report.applied_changes == [(VRM_A, "nominal_voltage", 0.7, 0.85)]

    apply_report(session, report)
    vrm = session.vrm(POWER_NET_A)
    assert vrm.sense_voltage == pytest.approx(0.85)
    assert not vrm.sense_override  # still auto, still following the row

    # An explicitly different sense in the workbook is honoured and pinned.
    _edit(book, SHEET_VRM, VRM_A, "Sense V", 0.8)
    apply_report(session, import_tables(session, book))
    assert session.vrm(POWER_NET_A).sense_voltage == pytest.approx(0.8)
    assert session.vrm(POWER_NET_A).sense_override


def test_a_value_equal_to_the_current_one_is_skipped(session: Session, book: Path) -> None:
    """No override flip for a cell the user merely retyped identically."""
    _edit(book, SHEET_VRM, VRM_A, "Nominal V", 0.7)
    _edit(book, SHEET_SINK, SINK_A, "Current (A)", 1.0)

    report = import_tables(session, book)
    assert report.applied_changes == []
    assert not session.vrm(POWER_NET_A).nominal_override


@pytest.mark.parametrize(
    ("cell", "expected"),
    [(False, False), (0, False), ("FALSE", False), ("no", False), (1, True), ("TRUE", True)],
)
def test_use_column_accepts_excel_booleans_and_one_zero(
    session: Session, book: Path, cell: Any, expected: bool
) -> None:
    _edit(book, SHEET_VRM, VRM_A, "Use", cell)
    report = import_tables(session, book)

    assert report.errors == []
    if expected:  # unchanged -- every row exports enabled
        assert report.applied_changes == []
    else:
        assert report.applied_changes == [(VRM_A, "enabled", True, False)]
        apply_report(session, report)
        assert session.vrm(POWER_NET_A).enabled is False


def test_component_and_ground_changes_apply(session: Session, book: Path) -> None:
    _edit(book, SHEET_SINK, SINK_A, "Component", "LGA")

    report = import_tables(session, book)
    assert report.errors == []
    assert report.applied_changes == [(SINK_A, "comp", "SITE0", "LGA")]

    apply_report(session, report)
    assert session.sink(POWER_NET_A).comp == "LGA"
    # `comp` has no dataclass override flag -- `Session._pinned` carries it.
    assert (SINK_A, "comp") in session._pinned


def test_row_order_in_the_workbook_does_not_matter(session: Session, book: Path) -> None:
    """The user sorted the sheet in Excel: same rows, reversed, still matched."""
    records = _rows(book, SHEET_VRM)
    records[0]["Nominal V"] = 0.9
    _rewrite(book, SHEET_VRM, list(reversed(records)))
    sinks = list(reversed(_rows(book, SHEET_SINK)))
    _rewrite(book, SHEET_SINK, sinks)

    assert _rows(book, SHEET_VRM)[0][KEY_HEADER] == VRM_B  # really reordered

    report = import_tables(session, book)
    assert report.errors == []
    assert report.warnings == []
    assert report.applied_changes == [(VRM_A, "nominal_voltage", 0.7, 0.9)]

    apply_report(session, report)
    assert session.vrm(POWER_NET_A).nominal_voltage == pytest.approx(0.9)
    assert session.vrm(POWER_NET_B).nominal_voltage == pytest.approx(1.2)


def test_column_order_in_the_workbook_does_not_matter(session: Session, book: Path) -> None:
    records = _rows(book, SHEET_VRM)
    records[0]["Sense V"] = 0.75
    shuffled = list(reversed(_headers(book, SHEET_VRM)))
    _rewrite(book, SHEET_VRM, records, headers=shuffled)

    report = import_tables(session, book)
    assert report.errors == []
    assert report.applied_changes == [(VRM_A, "sense_voltage", 0.7, 0.75)]


# --------------------------------------------------------------------------- #
# Tamper matrix -- every entry blocks the whole import
# --------------------------------------------------------------------------- #


def test_renamed_net_cell_is_blocked(session: Session, book: Path) -> None:
    before = _snapshot(session)
    _edit(book, SHEET_VRM, VRM_A, "Nominal V", 0.9)  # a legitimate edit alongside
    _edit(book, SHEET_VRM, VRM_A, "Net", "SOME_OTHER_RAIL/0")

    report = import_tables(session, book)
    _assert_blocked(session, report, "VRMs!C", "must not be edited", "SOME_OTHER_RAIL/0")
    assert _snapshot(session) == before  # the good edit did not sneak through


def test_edited_key_cell_reads_as_an_unknown_row(session: Session, book: Path) -> None:
    _edit(book, SHEET_SINK, SINK_A, KEY_HEADER, "sink:NOT_A_NET/0#0")

    report = import_tables(session, book)
    _assert_blocked(session, report, "Sinks!A", "not a row of the loaded design")


def test_a_key_from_the_other_sheet_is_blocked(session: Session, book: Path) -> None:
    _edit(book, SHEET_VRM, VRM_A, KEY_HEADER, SINK_A)

    report = import_tables(session, book)
    _assert_blocked(session, report, "does not belong on the 'VRMs' sheet")


def test_a_missing_key_cell_is_blocked(session: Session, book: Path) -> None:
    _edit(book, SHEET_VRM, VRM_A, KEY_HEADER, None)

    report = import_tables(session, book)
    _assert_blocked(session, report, "VRMs!A", "cannot tell which row it is")


def test_a_deleted_required_column_is_blocked(session: Session, book: Path) -> None:
    headers = [header for header in _headers(book, SHEET_SINK) if header != "Current (A)"]
    _rewrite(book, SHEET_SINK, _rows(book, SHEET_SINK), headers=headers)

    report = import_tables(session, book)
    _assert_blocked(session, report, "Current (A)", "column structure must not change")


def test_a_deleted_key_column_is_blocked(session: Session, book: Path) -> None:
    headers = [header for header in _headers(book, SHEET_VRM) if header != KEY_HEADER]
    _rewrite(book, SHEET_VRM, _rows(book, SHEET_VRM), headers=headers)

    report = import_tables(session, book)
    _assert_blocked(session, report, KEY_HEADER, "column structure must not change")


def test_a_duplicated_key_is_blocked(session: Session, book: Path) -> None:
    records = _rows(book, SHEET_VRM)
    clone = dict(records[0])
    clone["Nominal V"] = 0.9
    _rewrite(book, SHEET_VRM, [*records, clone])

    report = import_tables(session, book)
    _assert_blocked(session, report, "duplicate Key", VRM_A)


def test_a_duplicated_column_is_blocked(session: Session, book: Path) -> None:
    headers = [*_headers(book, SHEET_VRM), "Nominal V"]
    _rewrite(book, SHEET_VRM, _rows(book, SHEET_VRM), headers=headers)

    report = import_tables(session, book)
    _assert_blocked(session, report, "appears more than once")


def test_a_missing_sheet_is_blocked(session: Session, book: Path) -> None:
    workbook = load_workbook(book)
    del workbook[SHEET_SINK]
    workbook.save(book)
    workbook.close()

    report = import_tables(session, book)
    _assert_blocked(session, report, "no 'Sinks' sheet")


@pytest.mark.parametrize(
    ("sheet", "row_key", "header", "value", "needle"),
    [
        (SHEET_VRM, VRM_A, "Nominal V", -0.7, "must be greater than 0"),
        (SHEET_VRM, VRM_A, "Nominal V", 0, "must be greater than 0"),
        (SHEET_VRM, VRM_A, "Sense V", "zero point seven", "must be a number"),
        (SHEET_VRM, VRM_A, "Nominal V", None, "must be a number"),
        (SHEET_SINK, SINK_A, "Current (A)", -3, "must not be negative"),
        (SHEET_SINK, SINK_A, "Model", 2.5, "must be a whole number"),
        (SHEET_SINK, SINK_A, "PFMode", "two", "must be a whole number"),
        (SHEET_VRM, VRM_A, "Use", "maybe", "TRUE/FALSE or 1/0"),
        (SHEET_VRM, VRM_A, "Component", "NOT_A_CIRCUIT", "not a circuit of the loaded design"),
        (SHEET_SINK, SINK_A, "Ground", "NOT_A_GROUND", "not a ground net"),
    ],
)
def test_bad_cell_values_are_blocked(
    session: Session,
    book: Path,
    sheet: str,
    row_key: str,
    header: str,
    value: Any,
    needle: str,
) -> None:
    before = _snapshot(session)
    _edit(book, sheet, row_key, header, value)

    report = import_tables(session, book)
    _assert_blocked(session, report, needle)
    assert _snapshot(session) == before


def test_every_bad_cell_is_reported_not_just_the_first(session: Session, book: Path) -> None:
    _edit(book, SHEET_VRM, VRM_A, "Nominal V", -1)
    _edit(book, SHEET_VRM, VRM_B, "Sense V", "oops")
    _edit(book, SHEET_SINK, SINK_A, "Model", "x")

    report = import_tables(session, book)
    assert len(report.errors) == 3
    assert any("VRMs!F2" in error or "VRMs!F3" in error for error in report.errors)


def test_meta_schema_mismatch_is_a_clear_error(session: Session, book: Path) -> None:
    workbook = load_workbook(book)
    meta = workbook[SHEET_META]
    for row in meta.iter_rows(min_row=1, max_col=1):
        if row[0].value == "schema":
            meta.cell(row=row[0].row, column=2, value=SCHEMA_VERSION + 1)
    workbook.save(book)
    workbook.close()

    report = import_tables(session, book)
    _assert_blocked(session, report, "schema", "export the tables again")
    assert len(report.errors) == 1  # nothing else is even attempted


def test_a_file_that_is_not_a_workbook_is_an_error(session: Session, tmp_path: Path) -> None:
    bogus = tmp_path / "not-really.xlsx"
    bogus.write_bytes(b"this is not a zip archive")

    report = import_tables(session, bogus)
    _assert_blocked(session, report, "not a readable .xlsx workbook")


# --------------------------------------------------------------------------- #
# Warnings -- reported, never blocking
# --------------------------------------------------------------------------- #


def test_a_row_missing_from_the_workbook_is_only_a_warning(session: Session, book: Path) -> None:
    records = [record for record in _rows(book, SHEET_VRM) if record[KEY_HEADER] != VRM_B]
    records[0]["Nominal V"] = 0.8
    _rewrite(book, SHEET_VRM, records)

    report = import_tables(session, book)
    assert report.errors == []
    assert any(VRM_B in warning and "left unchanged" in warning for warning in report.warnings)
    assert report.applied_changes == [(VRM_A, "nominal_voltage", 0.7, 0.8)]

    apply_report(session, report)
    assert session.vrm(POWER_NET_B).nominal_voltage == pytest.approx(1.2)


def test_edited_read_only_cells_are_warned_about_and_ignored(session: Session, book: Path) -> None:
    _edit(book, SHEET_VRM, VRM_A, "Name", "MY_OWN_BLOCK_NAME")
    _edit(book, SHEET_VRM, VRM_A, "Pos pins", 999)

    report = import_tables(session, book)
    assert report.errors == []
    assert report.applied_changes == []
    assert sum("read-only" in warning for warning in report.warnings) == 2


def test_extra_columns_and_sheets_are_warned_about_and_ignored(
    session: Session, book: Path
) -> None:
    records = _rows(book, SHEET_VRM)
    for record in records:
        record["My notes"] = "check with layout"
    _rewrite(book, SHEET_VRM, records, headers=[*_headers(book, SHEET_VRM), "My notes"])
    workbook = load_workbook(book)
    workbook.create_sheet("Scratch").append(["free-form", "notes"])
    workbook.save(book)
    workbook.close()

    report = import_tables(session, book)
    assert report.errors == []
    assert report.applied_changes == []
    joined = " | ".join(report.warnings)
    assert "My notes" in joined and "Scratch" in joined


def test_a_workbook_without_meta_still_imports(session: Session, book: Path) -> None:
    workbook = load_workbook(book)
    del workbook[SHEET_META]
    workbook.save(book)
    workbook.close()

    report = import_tables(session, book)
    assert report.errors == []
    assert any("_meta" in warning for warning in report.warnings)


# --------------------------------------------------------------------------- #
# Applying
# --------------------------------------------------------------------------- #


def test_apply_report_refuses_a_report_that_carries_errors(session: Session) -> None:
    report = ImportReport(applied_changes=[(VRM_A, "nominal_voltage", 0.7, 0.9)], errors=["bad"])
    with pytest.raises(ValueError):
        apply_report(session, report)
    assert session.vrm(POWER_NET_A).nominal_voltage == pytest.approx(0.7)


def test_import_never_mutates_the_session_on_its_own(session: Session, book: Path) -> None:
    _edit(book, SHEET_VRM, VRM_A, "Nominal V", 0.95)
    before = _snapshot(session)

    report = import_tables(session, book)
    assert report.applied_changes  # it *would* change something
    assert _snapshot(session) == before  # ... but only when the caller says so

    apply_report(session, report)
    assert _snapshot(session) != before


def test_a_full_edit_pass_survives_export_import_export(session: Session, book: Path) -> None:
    """Edit every editable column of one row, then re-export: same values back."""
    _edit(book, SHEET_VRM, VRM_A, "Use", False)
    _edit(book, SHEET_VRM, VRM_A, "Component", "SITE0")
    _edit(book, SHEET_VRM, VRM_A, "Nominal V", 0.66)
    _edit(book, SHEET_VRM, VRM_A, "Sense V", 0.65)
    _edit(book, SHEET_VRM, VRM_A, "Output Current (A)", 42.5)

    report = import_tables(session, book)
    assert report.errors == []
    apply_report(session, report)

    again = export_tables(session, book.with_name("again.xlsx"))
    row = {record[KEY_HEADER]: record for record in _rows(again, SHEET_VRM)}[VRM_A]
    assert row["Use"] is False
    assert row["Component"] == "SITE0"
    assert (row["Nominal V"], row["Sense V"]) == (pytest.approx(0.66), pytest.approx(0.65))
    assert row["Output Current (A)"] == pytest.approx(42.5)
    assert import_tables(session, again).applied_changes == []
